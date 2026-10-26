from flask import render_template, request, redirect, url_for, flash, make_response
from flask_login import login_user, logout_user, login_required, current_user
from werkzeug.utils import secure_filename
from urllib.parse import quote
from io import BytesIO
from datetime import datetime as dt
import os
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app import app, db
from models import Sneaker, Size, Inventory
from auth import get_user

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}

NOMBRE_TIENDA = "JARVIS ALL SPORT"
WHATSAPP_NUMERO = "51925339422"
WHATSAPP_NUMERO_2 = "51956877923"
MARCAS = ["Nike", "Adidas", "Puma", "Lacoste", "Jordan", "Fila", "Reebok", "New Balance"]


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def guardar_imagen(imagen):
    if imagen and imagen.filename and allowed_file(imagen.filename):
        filename = secure_filename(imagen.filename)
        filename = f"{uuid.uuid4().hex}_{filename}"
        ruta = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        imagen.save(ruta)
        return f"uploads/{filename}"
    return None


def borrar_imagen(imagen_url):
    if not imagen_url:
        return
    try:
        nombre = imagen_url.split('/')[-1]
        ruta = os.path.join(app.config['UPLOAD_FOLDER'], nombre)
        if os.path.exists(ruta):
            os.remove(ruta)
    except Exception:
        pass


# ============================================================
# RUTAS PÚBLICAS (Cualquiera puede ver)
# ============================================================
@app.route('/')
def bienvenida():
    return render_template('bienvenida.html', nombre_tienda=NOMBRE_TIENDA)


@app.route('/catalogo')
def catalogo():
    marca_filtro = request.args.get('marca', '')
    categoria_filtro = request.args.get('categoria', '')

    query = Sneaker.query
    if marca_filtro:
        query = query.filter(Sneaker.marca == marca_filtro)
    if categoria_filtro:
        query = query.filter(Sneaker.categoria == categoria_filtro)

    sneakers = query.order_by(Sneaker.fecha_creacion.desc()).all()

    return render_template(
        'catalogo.html',
        sneakers=sneakers,
        marcas=MARCAS,
        marca_activa=marca_filtro,
        categoria_activa=categoria_filtro,
        nombre_tienda=NOMBRE_TIENDA,
        whatsapp=WHATSAPP_NUMERO
    )


@app.route('/detalle/<int:sneaker_id>')
def detalle(sneaker_id):
    sneaker = Sneaker.query.get_or_404(sneaker_id)
    inventarios = Inventory.query.join(Size).filter(
        Inventory.sneaker_id == sneaker_id
    ).order_by(Size.numero).all()

    mensaje = f"Hola {NOMBRE_TIENDA}, me interesa la zapatilla {sneaker.marca} {sneaker.nombre}"
    if sneaker.codigo:
        mensaje += f" (Cod. {sneaker.codigo})"
    mensaje += ". Esta disponible?"
    mensaje_url = quote(mensaje)

    return render_template(
        'detalle.html',
        sneaker=sneaker,
        inventarios=inventarios,
        nombre_tienda=NOMBRE_TIENDA,
        whatsapp=WHATSAPP_NUMERO,
        mensaje_url=mensaje_url
    )


@app.route('/compartir-catalogo')
def compartir_catalogo():
    sneakers = Sneaker.query.order_by(Sneaker.marca, Sneaker.nombre).all()

    lineas = [f"CATALOGO {NOMBRE_TIENDA}\n"]
    marca_actual = None

    for s in sneakers:
        if s.marca != marca_actual:
            marca_actual = s.marca
            lineas.append(f"\n=== {marca_actual.upper()} ===")

        linea = f"\n{s.nombre}"
        if s.codigo:
            linea += f"\n   Cod: {s.codigo}"
        linea += f"\n   Precio: S/ {s.precio:.2f}"

        tallas = [inv.size.numero for inv in s.inventarios]
        if tallas:
            linea += f"\n   Tallas: {', '.join(tallas)}"

        lineas.append(linea)

    lineas.append(f"\n\nPide el tuyo al WhatsApp: +{WHATSAPP_NUMERO}")

    mensaje = "\n".join(lineas)
    mensaje_url = quote(mensaje)

    return redirect(f"https://wa.me/?text={mensaje_url}")


# ============================================================
# RUTAS DE SEGURIDAD (LOGIN Y LOGOUT)
# ============================================================
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = get_user(username)
        
        # Verificamos si el usuario existe y la contraseña es correcta
        if user and user.check_password(password):
            login_user(user)
            flash('Bienvenido al panel de administración', 'success')
            return redirect(url_for('admin_dashboard'))
        else:
            flash('Usuario o contraseña incorrectos', 'error')
            
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('Has cerrado sesión', 'success')
    return redirect(url_for('catalogo'))


# ============================================================
# RUTAS DE ADMINISTRACIÓN (Protegidas con contraseña)
# ============================================================
@app.route('/admin')
@login_required  # <--- Esto protege la ruta
def admin_dashboard():
    marca_filtro = request.args.get('marca', '')
    query = Sneaker.query
    if marca_filtro:
        query = query.filter(Sneaker.marca == marca_filtro)
    sneakers = query.order_by(Sneaker.fecha_creacion.desc()).all()
    return render_template('admin_dashboard.html', sneakers=sneakers, marcas=MARCAS, marca_activa=marca_filtro)


@app.route('/admin/agregar', methods=['GET', 'POST'])
@login_required  # <--- Esto protege la ruta
def agregar_sneaker():
    if request.method == 'POST':
        nombre = request.form.get('nombre', '').strip()
        marca = request.form.get('marca', '').strip()
        genero = request.form.get('genero', 'Caballero')
        categoria = request.form.get('categoria', 'Adulto')
        codigo = request.form.get('codigo', '').strip()
        descripcion = request.form.get('descripcion', '').strip()
        try:
            precio = float(request.form.get('precio', 0))
        except ValueError:
            flash('Precio invalido', 'error')
            return redirect(url_for('agregar_sneaker'))

        precio_oferta_str = request.form.get('precio_oferta', '').strip()
        precio_oferta = float(precio_oferta_str) if precio_oferta_str else None

        imagen_url = guardar_imagen(request.files.get('imagen'))

        nueva = Sneaker(
            nombre=nombre, marca=marca, genero=genero, categoria=categoria,
            codigo=codigo, descripcion=descripcion, precio=precio,
            precio_oferta=precio_oferta, imagen_url=imagen_url
        )
        db.session.add(nueva)
        db.session.flush()

        tallas = request.form.getlist('talla[]')
        cantidades = request.form.getlist('cantidad[]')
        for t, c in zip(tallas, cantidades):
            t = t.strip()
            if not t or not c:
                continue
            try:
                cant = int(c)
            except ValueError:
                continue
            size = Size.query.filter_by(numero=t).first()
            if not size:
                size = Size(numero=t)
                db.session.add(size)
                db.session.flush()
            db.session.add(Inventory(sneaker_id=nueva.id, size_id=size.id, cantidad=cant))

        db.session.commit()
        flash('Zapatilla agregada', 'success')
        return redirect(url_for('admin_dashboard'))

    return render_template('admin_form.html', sneaker=None, inventarios=[], marcas=MARCAS)


@app.route('/admin/editar/<int:sneaker_id>', methods=['GET', 'POST'])
@login_required  # <--- Esto protege la ruta
def editar_sneaker(sneaker_id):
    sneaker = Sneaker.query.get_or_404(sneaker_id)
    if request.method == 'POST':
        sneaker.nombre = request.form.get('nombre', '').strip()
        sneaker.marca = request.form.get('marca', '').strip()
        sneaker.genero = request.form.get('genero', 'Caballero')
        sneaker.categoria = request.form.get('categoria', 'Adulto')
        sneaker.codigo = request.form.get('codigo', '').strip()
        sneaker.descripcion = request.form.get('descripcion', '').strip()
        try:
            sneaker.precio = float(request.form.get('precio', 0))
        except ValueError:
            flash('Precio invalido', 'error')
            return redirect(url_for('editar_sneaker', sneaker_id=sneaker.id))

        precio_oferta_str = request.form.get('precio_oferta', '').strip()
        sneaker.precio_oferta = float(precio_oferta_str) if precio_oferta_str else None

        imagen = request.files.get('imagen')
        if imagen and imagen.filename:
            nueva_url = guardar_imagen(imagen)
            if nueva_url:
                borrar_imagen(sneaker.imagen_url)
                sneaker.imagen_url = nueva_url

        Inventory.query.filter_by(sneaker_id=sneaker.id).delete()
        tallas = request.form.getlist('talla[]')
        cantidades = request.form.getlist('cantidad[]')
        for t, c in zip(tallas, cantidades):
            t = t.strip()
            if not t or not c:
                continue
            try:
                cant = int(c)
            except ValueError:
                continue
            size = Size.query.filter_by(numero=t).first()
            if not size:
                size = Size(numero=t)
                db.session.add(size)
                db.session.flush()
            db.session.add(Inventory(sneaker_id=sneaker.id, size_id=size.id, cantidad=cant))

        db.session.commit()
        flash('Zapatilla actualizada', 'success')
        return redirect(url_for('admin_dashboard'))

    inventarios = Inventory.query.join(Size).filter(
        Inventory.sneaker_id == sneaker.id
    ).order_by(Size.numero).all()
    return render_template('admin_form.html', sneaker=sneaker, inventarios=inventarios, marcas=MARCAS)


@app.route('/admin/eliminar/<int:sneaker_id>', methods=['POST'])
@login_required  # <--- Esto protege la ruta
def eliminar_sneaker(sneaker_id):
    sneaker = Sneaker.query.get_or_404(sneaker_id)
    borrar_imagen(sneaker.imagen_url)
    db.session.delete(sneaker)
    db.session.commit()
    flash('Zapatilla eliminada', 'warning')
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/excel')
@login_required  # <--- Esto protege la ruta
def descargar_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    fuente_encabezado = Font(bold=True, color="FFFFFF", size=12)
    fondo_encabezado = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    centrado = Alignment(horizontal="center", vertical="center")

    encabezados = ["ID", "Marca", "Nombre", "Codigo", "Genero", "Categoria",
                   "Precio", "Precio Oferta", "Tallas", "Stock Total", "Fecha"]
    ws.append(encabezados)

    for col_num, _ in enumerate(encabezados, 1):
        celda = ws.cell(row=1, column=col_num)
        celda.font = fuente_encabezado
        celda.fill = fondo_encabezado
        celda.alignment = centrado

    sneakers = Sneaker.query.order_by(Sneaker.marca, Sneaker.nombre).all()
    for s in sneakers:
        tallas = ", ".join([inv.size.numero for inv in s.inventarios])
        stock_total = sum(inv.cantidad for inv in s.inventarios)
        fecha = s.fecha_creacion.strftime("%d/%m/%Y %H:%M") if s.fecha_creacion else ""

        ws.append([
            s.id, s.marca or "", s.nombre or "", s.codigo or "",
            s.genero or "", s.categoria or "", s.precio,
            s.precio_oferta if s.precio_oferta else "",
            tallas, stock_total, fecha
        ])

    anchos = [6, 12, 30, 18, 12, 12, 10, 12, 30, 10, 18]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename=productos-{dt.now().strftime("%Y%m%d")}.xlsx'
    return response


@app.route('/admin/catalogo-pdf')
@login_required  # <--- Esto protege la ruta
def catalogo_pdf():
    from fpdf import FPDF

    sneakers = Sneaker.query.order_by(Sneaker.marca, Sneaker.nombre).all()

    pdf = FPDF()
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(22, 163, 74)
    pdf.cell(0, 12, "CATALOGO " + NOMBRE_TIENDA, ln=True, align="C")

    pdf.set_font("Helvetica", "I", 11)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 8, "La tienda #1 en toda la selva", ln=True, align="C")
    pdf.ln(5)

    marcas_dict = {}
    for s in sneakers:
        marcas_dict.setdefault(s.marca or 'Otras', []).append(s)

    for marca, productos in marcas_dict.items():
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(255, 255, 255)
        pdf.set_fill_color(15, 23, 42)
        pdf.cell(0, 10, "  " + marca.upper(), ln=True, fill=True)
        pdf.ln(2)

        for s in productos:
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(0, 0, 0)
            pdf.cell(0, 7, s.nombre, ln=True)

            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(100, 100, 100)
            info = f"Codigo: {s.codigo}" if s.codigo else ""
            if s.genero:
                info += f"   Genero: {s.genero}"
            if info:
                pdf.cell(0, 5, info, ln=True)

            tallas = ", ".join([inv.size.numero for inv in s.inventarios])
            if tallas:
                pdf.set_font("Helvetica", "", 9)
                pdf.set_text_color(60, 60, 60)
                pdf.cell(0, 5, "Tallas: " + tallas, ln=True)

            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(22, 163, 74)
            if s.precio_oferta:
                pdf.cell(0, 7, f"S/ {s.precio:.2f}   (antes S/ {s.precio_oferta:.2f})", ln=True)
            else:
                pdf.cell(0, 7, f"S/ {s.precio:.2f}", ln=True)

            pdf.ln(3)

    pdf.ln(5)
    pdf.set_font("Helvetica", "I", 10)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 8, f"Pide el tuyo al WhatsApp: +{WHATSAPP_NUMERO}", ln=True, align="C")

    output = pdf.output(dest='S').encode('latin-1')
    response = make_response(output)
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'attachment; filename=catalogo-{dt.now().strftime("%Y%m%d")}.pdf'
    return response